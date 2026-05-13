#!/usr/bin/env python3
"""
Generate Lab 6 Word document: "Работа с формулами" (Variant 1)
Contains formulas from Tables 6.1-6.5 inserted via OMML (Office Math Markup Language).
"""

from docx import Document
from docx.shared import Cm, Pt, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree
import copy

OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
WPC_NS = "http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

nsmap = {
    "m": OMML_NS,
    "w": W_NS,
}


def make_omml(xml_str: str):
    """Parse an OMML XML string and return the element."""
    full = f'<m:oMath xmlns:m="{OMML_NS}" xmlns:w="{W_NS}">{xml_str}</m:oMath>'
    return etree.fromstring(full.encode())


def run_text(text, italic=True, sz=28):
    """Create an <m:r> element with the given text."""
    style = "p" if italic else ""
    return f"""<m:r><m:rPr><m:sty m:val="{'p' if not italic else 'i'}"/></m:rPr><w:rPr><w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr><m:t>{text}</m:t></m:r>"""


def r(text, italic=True):
    """Shortcut for run_text with default 14pt (28 half-points)."""
    return run_text(text, italic=italic, sz=28)


# ──────────────────────────────────────────────
# Table 6.1: x^3 + 3x^2 - 2 = 0
# ──────────────────────────────────────────────
EQ1_OMML = f"""
<m:sSup><m:sSupPr><m:ctrlPr/></m:sSupPr><m:e>{r("x")}</m:e><m:sup>{r("3")}</m:sup></m:sSup>
{r("+")}
{r("3")}
<m:sSup><m:sSupPr><m:ctrlPr/></m:sSupPr><m:e>{r("x")}</m:e><m:sup>{r("2")}</m:sup></m:sSup>
{r("−")}
{r("2")}
{r("=")}
{r("0")}
"""

# ──────────────────────────────────────────────
# Table 6.2: System of equations
# sin(x+1) - y = 1.2
# 2x + cos(y) = 2
# ──────────────────────────────────────────────
SYS_LINE1 = f"""{r("sin")}
<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{r("x")}{r("+")}{r("1")}</m:e></m:d>
{r("−")}{r("y")}{r("=")}{r("1.2;")}"""

SYS_LINE2 = f"""{r("2")}{r("x")}{r("+")}{r("cos")}
<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{r("y")}</m:e></m:d>
{r("=")}{r("2.")}"""

EQ2_OMML = f"""
<m:d>
  <m:dPr>
    <m:begChr m:val="{{"/>
    <m:endChr m:val=""/>
  </m:dPr>
  <m:e>
    <m:eqArr>
      <m:e>{SYS_LINE1}</m:e>
      <m:e>{SYS_LINE2}</m:e>
    </m:eqArr>
  </m:e>
</m:d>
"""

# ──────────────────────────────────────────────
# Table 6.3: Integral ∫₀¹ (sin x)/(1+x²) dx
# ──────────────────────────────────────────────
EQ3_OMML = f"""
<m:nary>
  <m:naryPr>
    <m:chr m:val="∫"/>
    <m:limLoc m:val="subSup"/>
  </m:naryPr>
  <m:sub>{r("0")}</m:sub>
  <m:sup>{r("1")}</m:sup>
  <m:e>
    <m:f>
      <m:fPr><m:type m:val="bar"/></m:fPr>
      <m:num>{r("sin")}{r(" x")}</m:num>
      <m:den>{r("1")}{r("+")}<m:sSup><m:sSupPr><m:ctrlPr/></m:sSupPr><m:e>{r("x")}</m:e><m:sup>{r("2")}</m:sup></m:sSup></m:den>
    </m:f>
    {r(" dx")}
  </m:e>
</m:nary>
"""

# ──────────────────────────────────────────────
# Table 6.4: y' = x + cos(y/√5), y(1.8)=2.6, [1.8;2.8]
# ──────────────────────────────────────────────
EQ4_OMML = f"""
{r("y′")}{r("=")}{r("x")}{r("+")}
{r("cos")}
<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>
<m:e>
  <m:f>
    <m:fPr><m:type m:val="bar"/></m:fPr>
    <m:num>{r("y")}</m:num>
    <m:den><m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{r("5")}</m:e></m:rad></m:den>
  </m:f>
</m:e>
</m:d>
{r(",")}
{r(" ")}
{r("y")}
<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>{r("1,8")}</m:e></m:d>
{r("=")}
{r("2,6")}
{r(",")}
{r(" ")}
<m:d><m:dPr><m:begChr m:val="["/><m:endChr m:val="]"/></m:dPr><m:e>{r("1,8;")}{r(" ")}{r("2,8")}</m:e></m:d>
"""

# ──────────────────────────────────────────────
# Table 6.5 (Variant 1,11):
# Matrix A (3x3) * vector b (3x1) + 2 * vector c (3x1)
# A = ((3, -2, -1), (4, -1, -3), (2, -1, -1))
# b = ((1), (-2), (5))
# c = ((4), (-2), (3))
# Expression: A·b + 2·c
# ──────────────────────────────────────────────

def matrix_3x3(rows):
    """Generate OMML for a 3x3 matrix."""
    row_xml = ""
    for row in rows:
        cols = "".join(f"<m:e>{r(str(v))}</m:e>" for v in row)
        row_xml += f"<m:mr>{cols}</m:mr>"
    return f"""<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>
    <m:m><m:mPr><m:mcs><m:mc><m:mcPr><m:count m:val="{len(rows[0])}"/><m:mcJc m:val="center"/></m:mcPr></m:mc></m:mcs></m:mPr>
    {row_xml}</m:m></m:e></m:d>"""


def col_vector(vals):
    """Generate OMML for a column vector."""
    row_xml = "".join(f"<m:mr><m:e>{r(str(v))}</m:e></m:mr>" for v in vals)
    return f"""<m:d><m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr><m:e>
    <m:m><m:mPr><m:mcs><m:mc><m:mcPr><m:count m:val="1"/><m:mcJc m:val="center"/></m:mcPr></m:mc></m:mcs></m:mPr>
    {row_xml}</m:m></m:e></m:d>"""


A = [[3, -2, -1], [4, -1, -3], [2, -1, -1]]
b = [1, -2, 5]
c = [4, -2, 3]

EQ5_OMML = f"""
{matrix_3x3(A)}
{r("·")}
{col_vector(b)}
{r("+")}
{r("2")}
{col_vector(c)}
"""


def add_formula_to_paragraph(paragraph, omml_xml_str):
    """Add an OMML formula to an existing paragraph."""
    omath = make_omml(omml_xml_str)
    paragraph._element.append(omath)


def create_lab6():
    doc = Document()

    # Set margins to 2 cm
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Title
    title = doc.add_heading("Работа с формулами", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph("Вариант 1")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].bold = True

    # ─── Table 6.1: Equation ───
    doc.add_heading("Таблица 6.1 – Уравнение", level=2)
    p1_label = doc.add_paragraph()
    p1_label.add_run("Способ 1 (редактор формул):").bold = True
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formula_to_paragraph(p1, EQ1_OMML)

    p1b_label = doc.add_paragraph()
    p1b_label.add_run("Способ 2 (линейный формат):").bold = True
    p1b = doc.add_paragraph("x³ + 3x² − 2 = 0")
    p1b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p1b.runs:
        run.font.size = Pt(14)

    doc.add_paragraph()  # spacer

    # ─── Table 6.2: System of equations ───
    doc.add_heading("Таблица 6.2 – Система уравнений", level=2)
    p2_label = doc.add_paragraph()
    p2_label.add_run("Способ 1 (редактор формул):").bold = True
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formula_to_paragraph(p2, EQ2_OMML)

    p2b_label = doc.add_paragraph()
    p2b_label.add_run("Способ 2 (линейный формат):").bold = True
    p2b = doc.add_paragraph("sin(x + 1) − y = 1,2;\n2x + cos(y) = 2.")
    p2b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p2b.runs:
        run.font.size = Pt(14)

    doc.add_paragraph()

    # ─── Table 6.3: Integral ───
    doc.add_heading("Таблица 6.3 – Интеграл", level=2)
    p3_label = doc.add_paragraph()
    p3_label.add_run("Способ 1 (редактор формул):").bold = True
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formula_to_paragraph(p3, EQ3_OMML)

    p3b_label = doc.add_paragraph()
    p3b_label.add_run("Способ 2 (линейный формат):").bold = True
    p3b = doc.add_paragraph("∫₀¹ (sin x)/(1 + x²) dx")
    p3b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p3b.runs:
        run.font.size = Pt(14)

    doc.add_paragraph()

    # ─── Table 6.4: Differential equation ───
    doc.add_heading("Таблица 6.4 – Дифференциальное уравнение", level=2)
    p4_label = doc.add_paragraph()
    p4_label.add_run("Способ 1 (редактор формул):").bold = True
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formula_to_paragraph(p4, EQ4_OMML)

    p4b_label = doc.add_paragraph()
    p4b_label.add_run("Способ 2 (линейный формат):").bold = True
    p4b = doc.add_paragraph("y′ = x + cos(y/√5),  y(1,8) = 2,6,  [1,8; 2,8]")
    p4b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p4b.runs:
        run.font.size = Pt(14)

    doc.add_paragraph()

    # ─── Table 6.5: Matrix computation ───
    doc.add_heading("Таблица 6.5 – Матричные вычисления", level=2)
    p5_label = doc.add_paragraph()
    p5_label.add_run("Способ 1 (редактор формул):").bold = True
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_formula_to_paragraph(p5, EQ5_OMML)

    p5b_label = doc.add_paragraph()
    p5b_label.add_run("Способ 2 (линейный формат):").bold = True
    p5b_text = (
        "( 3  −2  −1 )       ( 1 )       ( 4 )\n"
        "( 4  −1  −3 )  ·  ( −2 )  + 2( −2 )\n"
        "( 2  −1  −1 )       ( 5 )       ( 3 )"
    )
    p5b = doc.add_paragraph(p5b_text)
    p5b.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p5b.runs:
        run.font.size = Pt(14)
        run.font.name = "Courier New"

    output_path = "/home/ubuntu/word/Работа с формулами.docx"
    doc.save(output_path)
    print(f"Saved: {output_path}")
    return output_path


if __name__ == "__main__":
    create_lab6()
